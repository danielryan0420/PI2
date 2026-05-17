import os
import csv
import io
import json
from pathlib import Path
from datetime import datetime
from flask import Flask, jsonify, request, send_from_directory, send_file, Response
from flask_cors import CORS
from werkzeug.utils import secure_filename
import uuid
import openpyxl
from services import (
    UserService, SessionService, CountService, PhotoService,
    MessageService, AuditService, SlocConfigService, MaterialService,
    WmBinService, DashboardService, ImportService, ThreadService,
    OrderValidationService
)

app = Flask(__name__)
CORS(app)

DIST_DIR = Path(__file__).parent / 'client' / 'dist'

@app.before_request
def _track_activity():
    username = request.headers.get('x-username')
    if username and request.path.startswith('/api/'):
        try:
            UserService.update_last_active(username)
        except Exception:
            pass

@app.after_request
def _security_headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'SAMEORIGIN'
    response.headers['X-XSS-Protection'] = '1; mode=block'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    response.headers['Permissions-Policy'] = 'geolocation=(), microphone=()'
    return response

UPLOAD_FOLDER = Path('uploads')
UPLOAD_FOLDER.mkdir(exist_ok=True)
ALLOWED_EXTENSIONS = {'jpg', 'jpeg', 'png'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

# ===== AUTH =====
@app.post('/api/login')
def login():
    data = request.json
    username = data.get('username')
    password = data.get('password') or ''

    if not username:
        return jsonify({'error': 'Username required'}), 400

    user = UserService.get_user(username)
    if not user or not UserService.verify_password(username, password):
        return jsonify({'error': 'Invalid username or password'}), 401

    UserService.update_last_active(username)
    return jsonify({
        'id': user['id'],
        'username': user['username'],
        'role': user['role']
    }), 200

# ===== USERS =====
@app.get('/api/users')
def list_users():
    users = UserService.list_users()
    return jsonify(users)

@app.post('/api/users')
def create_user():
    data = request.json
    username = data.get('username')
    role = data.get('role')
    password = data.get('password', '')

    if not username or not role:
        return jsonify({'error': 'Username and role required'}), 400

    try:
        user = UserService.get_user(username)
        if user:
            UserService.update_user_role(username, role)
            if password:
                UserService.set_password(username, password)
            return jsonify({'id': user['id'], 'username': user['username'], 'role': user['role']})
        else:
            UserService.create_user(username, role, password)
            return jsonify(UserService.get_user(username)), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/users/<username>')
def get_user(username):
    user = UserService.get_user(username)
    if not user:
        return jsonify({'error': 'User not found'}), 404
    return jsonify({'id': user['id'], 'username': user['username'], 'role': user['role'], 'created_at': user['created_at']})

@app.patch('/api/users/<int:user_id>')
def update_user(user_id):
    data = request.json or {}
    role = data.get('role')
    if not role:
        return jsonify({'error': 'role required'}), 400
    try:
        UserService.update_user_role_by_id(user_id, role)
        return jsonify({'ok': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.delete('/api/users/<int:user_id>')
def delete_user(user_id):
    try:
        UserService.delete_user(user_id)
        return '', 204
    except Exception as e:
        return jsonify({'error': str(e)}), 400

# ===== SESSIONS =====
@app.get('/api/sessions')
def list_sessions():
    sessions = SessionService.list_sessions()
    return jsonify(sessions)

@app.post('/api/sessions')
def create_session():
    data = request.json
    name = data.get('name')
    if not name:
        return jsonify({'error': 'Session name required'}), 400

    try:
        session_id = SessionService.create_session(name)
        return jsonify(SessionService.get_session(session_id)), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sessions/<int:session_id>')
def get_session(session_id):
    session = SessionService.get_session(session_id)
    if not session:
        return jsonify({'error': 'Session not found'}), 404
    return jsonify(session)

@app.route('/api/sessions/<int:session_id>/close', methods=['POST', 'PATCH'])
def close_session(session_id):
    try:
        SessionService.close_session(session_id)
        return jsonify(SessionService.get_session(session_id))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sessions/<int:session_id>/stats')
def get_session_stats(session_id):
    stats = SessionService.get_session_stats(session_id)
    return jsonify(stats)

# ===== COUNTS =====
@app.post('/api/sessions/<int:session_id>/counts')
def create_count(session_id):
    data = request.json

    try:
        validation_warnings_raw = data.get('validation_warnings')
        count_id = CountService.create_count(
            session_id=session_id,
            username=data.get('username'),
            material_number=data.get('material_number'),
            quantity=float(data.get('quantity')),
            sloc=data.get('sloc'),
            wm_bin=data.get('wm_bin'),
            zbin=data.get('zbin'),
            validation_warnings=validation_warnings_raw
        )

        # Auto-post system message so admins are notified via the messages inbox
        if validation_warnings_raw:
            try:
                warnings = json.loads(validation_warnings_raw)
                mat = (data.get('material_number') or '').upper()
                wm = (data.get('wm_bin') or '').upper()
                counter = data.get('username', 'unknown')
                lines = []

                if 'material_not_found' in warnings:
                    lines.append(
                        f"Material {mat} was not found in Material Master (MARA). "
                        "Verify this is a valid material number before verifying the count."
                    )
                if 'wm_bin_not_found' in warnings:
                    lines.append(
                        f"WM Bin {wm} was not found in storage bin master (LGAP). "
                        "Verify this bin exists in the warehouse system."
                    )
                if 'fixed_bin_mismatch' in warnings and wm:
                    check = WmBinService.check_fixed_bin(mat, wm)
                    expected = check.get('expected_bin') or 'unknown'
                    lines.append(
                        f"Fixed bin mismatch: {mat} has fixed bin {expected} in MM02 WM2, "
                        f"but was counted into {wm}. Confirm the correct bin with the counter."
                    )

                if lines:
                    body = (
                        f"⚠️ Validation alert on Count #{count_id} submitted by {counter}:\n\n"
                        + "\n\n".join(f"• {l}" for l in lines)
                    )
                    MessageService.create_message(
                        session_id=session_id,
                        sender='SYSTEM',
                        role='counter',
                        body=body,
                        count_id=count_id
                    )
            except Exception:
                pass  # never let message creation block the count response

        return jsonify(CountService.get_count(count_id)), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sessions/<int:session_id>/counts')
def list_counts(session_id):
    material = request.args.get('material')
    status = request.args.get('status')
    username = request.args.get('username')
    sloc = request.args.get('sloc')

    filters = {}
    if material:
        filters['material'] = material
    if status:
        filters['status'] = status
    if username:
        filters['username'] = username
    if sloc:
        filters['sloc'] = sloc

    counts = CountService.list_counts(session_id, filters)
    return jsonify(counts)

@app.get('/api/sessions/<int:session_id>/counts/mine')
def list_my_counts(session_id):
    username = request.args.get('username') or request.headers.get('x-username')
    if not username:
        return jsonify({'error': 'username required'}), 400

    counts = CountService.list_counts(session_id, {'username': username})
    return jsonify(counts)

@app.get('/api/counts/<int:count_id>')
def get_count(count_id):
    count = CountService.get_count(count_id)
    if not count:
        return jsonify({'error': 'Count not found'}), 404
    return jsonify(count)

@app.patch('/api/counts/<int:count_id>')
def update_count(count_id):
    data = request.json
    # Accept both editor_username and editedBy (sent by OfficePage)
    editor_username = data.get('editor_username') or data.get('editedBy') or request.headers.get('x-username')
    reason = data.get('reason')
    # Accept either flat fields or nested {changes: {...}}
    changes = data.get('changes') or {k: v for k, v in data.items() if k not in ('editor_username', 'editedBy', 'reason', 'changes')}
    try:
        CountService.update_count(count_id, editor_username, reason=reason, **changes)
        return jsonify(CountService.get_count(count_id))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/api/counts/<int:count_id>/verify', methods=['POST', 'PATCH'])
def verify_count(count_id):
    data = request.json or {}
    editor_username = data.get('editor_username') or data.get('verifiedBy') or request.headers.get('x-username')
    try:
        CountService.verify_count(count_id, editor_username)
        return jsonify(CountService.get_count(count_id))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/api/counts/<int:count_id>/flag', methods=['POST', 'PATCH'])
def flag_count(count_id):
    data = request.json or {}
    editor_username = data.get('editor_username') or data.get('flaggedBy') or request.headers.get('x-username')
    reason = data.get('reason', '')
    try:
        CountService.flag_count(count_id, editor_username, reason)
        return jsonify(CountService.get_count(count_id))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.delete('/api/counts/<int:count_id>')
def delete_count(count_id):
    editor_username = request.headers.get('x-username') or ''
    count = CountService.get_count(count_id)
    if not count:
        return jsonify({'error': 'Count not found'}), 404
    try:
        CountService.delete_count(count_id, editor_username)
        return jsonify({'deleted': count_id}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 400

# ===== PHOTOS =====
@app.post('/api/counts/<int:count_id>/photos')
def upload_photo(count_id):
    if 'photo' not in request.files:
        return jsonify({'error': 'No photo provided'}), 400

    file = request.files['photo']
    if not file or file.filename == '':
        return jsonify({'error': 'No file selected'}), 400

    if not allowed_file(file.filename):
        return jsonify({'error': 'Invalid file type'}), 400

    try:
        ext = file.filename.rsplit('.', 1)[1].lower()
        filename = f"{count_id}_{uuid.uuid4()}.{ext}"
        filepath = UPLOAD_FOLDER / filename
        file.save(filepath)

        photo_id = PhotoService.add_photo(count_id, filename, file.filename)
        return jsonify({'id': photo_id, 'filename': filename}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/counts/<int:count_id>/photos')
def get_photos(count_id):
    photos = PhotoService.get_photos(count_id)
    return jsonify(photos)

@app.delete('/api/photos/<int:photo_id>')
def delete_photo(photo_id):
    try:
        PhotoService.delete_photo(photo_id)
        return '', 204
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/photos/<filename>')
def serve_photo(filename):
    try:
        return send_from_directory(UPLOAD_FOLDER, filename)
    except Exception as e:
        return jsonify({'error': 'Photo not found'}), 404

# ===== MESSAGES =====
@app.post('/api/sessions/<int:session_id>/messages')
def create_message(session_id):
    data = request.json

    # Normalize role — old cached sessions may have 'office', map to 'admin'
    raw_role = data.get('role') or request.headers.get('x-role') or 'counter'
    role = 'admin' if raw_role not in ('counter', 'admin') else raw_role

    try:
        message_id = MessageService.create_message(
            session_id=session_id,
            sender=data.get('sender'),
            role=role,
            body=data.get('body'),
            count_id=data.get('count_id')
        )
        return jsonify({'id': message_id}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sessions/<int:session_id>/messages')
def get_messages(session_id):
    count_id = request.args.get('count_id', type=int)
    messages = MessageService.get_messages(session_id, count_id)
    return jsonify(messages)

@app.get('/api/sessions/<int:session_id>/messages/mine')
def get_my_messages(session_id):
    username = request.args.get('username') or request.headers.get('x-username')
    if not username:
        return jsonify({'error': 'username required'}), 400
    messages = MessageService.get_user_messages(session_id, username)
    return jsonify(messages)

@app.get('/api/sessions/<int:session_id>/messages/general')
def get_general_messages(session_id):
    messages = MessageService.get_general_messages(session_id)
    return jsonify(messages)

@app.delete('/api/messages/<int:message_id>')
def delete_message(message_id):
    deleted = MessageService.delete_message(message_id)
    if not deleted:
        return jsonify({'error': 'Message not found'}), 404
    return jsonify({'deleted': True})

@app.get('/api/counts/<int:count_id>/messages')
def get_count_messages(count_id):
    messages = MessageService.get_count_messages(count_id)
    return jsonify(messages)

@app.post('/api/counts/<int:count_id>/messages')
def create_count_message(count_id):
    data = request.json
    count = CountService.get_count(count_id)
    if not count:
        return jsonify({'error': 'Count not found'}), 404
    raw_role = data.get('role') or request.headers.get('x-role') or 'counter'
    role = 'admin' if raw_role not in ('counter', 'admin') else raw_role
    try:
        message_id = MessageService.create_message(
            session_id=count['session_id'],
            sender=data.get('sender'),
            role=role,
            body=data.get('body'),
            count_id=count_id
        )
        return jsonify({'id': message_id}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

# ===== THREADS =====
@app.post('/api/sessions/<int:session_id>/threads')
def create_thread(session_id):
    data = request.json or {}
    title = (data.get('title') or '').strip()
    sender = request.headers.get('x-username', '')
    role = request.headers.get('x-role', 'counter')
    count_id = data.get('count_id')

    if not title or not sender:
        return jsonify({'error': 'title and x-username header required'}), 400

    session = SessionService.get_session(session_id)
    if not session:
        return jsonify({'error': 'Session not found'}), 404

    thread_role = 'admin' if role not in ('counter', 'admin') else role
    thread_id = ThreadService.create_thread(session_id, title, sender, thread_role, count_id)
    thread = ThreadService.get_thread(thread_id)
    return jsonify(thread), 201

@app.get('/api/sessions/<int:session_id>/threads')
def list_threads(session_id):
    username = request.headers.get('x-username', '')
    role = request.headers.get('x-role', 'counter')
    if not username:
        return jsonify({'error': 'x-username header required'}), 400
    thread_role = 'admin' if role not in ('counter', 'admin') else role
    threads = ThreadService.get_threads(session_id, username, thread_role)
    return jsonify(threads)

@app.get('/api/threads/<int:thread_id>/messages')
def get_thread_messages(thread_id):
    thread = ThreadService.get_thread(thread_id)
    if not thread:
        return jsonify({'error': 'Thread not found'}), 404
    messages = ThreadService.get_thread_messages(thread_id)
    return jsonify(messages)

@app.post('/api/threads/<int:thread_id>/messages')
def post_thread_message(thread_id):
    data = request.json or {}
    body = (data.get('body') or '').strip()
    sender = request.headers.get('x-username', '')
    role = request.headers.get('x-role', 'counter')
    reply_to_id = data.get('reply_to_id')

    if not body or not sender:
        return jsonify({'error': 'body and x-username header required'}), 400

    thread_role = 'admin' if role not in ('counter', 'admin') else role
    try:
        message = ThreadService.post_message(thread_id, sender, thread_role, body, reply_to_id)
        return jsonify(message), 201
    except ValueError as e:
        return jsonify({'error': str(e)}), 404

# ===== SLOC CONFIG =====
@app.get('/api/sloc-config')
def list_sloc_config():
    slocs = SlocConfigService.list_slocs()
    return jsonify(slocs)

@app.get('/api/sloc-config/available')
def list_available_slocs():
    result = SlocConfigService.get_available_slocs()
    return jsonify(result)

@app.post('/api/sloc-config')
def create_sloc_config():
    data = request.json

    try:
        SlocConfigService.create_sloc(
            sloc=data.get('sloc'),
            description=data.get('description'),
            wm_enabled=data.get('wm_enabled', False),
            im_enabled=data.get('im_enabled', False)
        )
        return jsonify(SlocConfigService.get_sloc(data.get('sloc'))), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sloc-config/<sloc>')
def get_sloc_config(sloc):
    config = SlocConfigService.get_sloc(sloc)
    if not config:
        return jsonify({'error': 'SLOC not found'}), 404
    return jsonify(config)

@app.patch('/api/sloc-config/<sloc>')
def update_sloc_config(sloc):
    data = request.json
    try:
        SlocConfigService.create_sloc(
            sloc=sloc,
            description=data.get('description', ''),
            wm_enabled=data.get('wm_enabled', False),
            im_enabled=data.get('im_enabled', False)
        )
        return jsonify(SlocConfigService.get_sloc(sloc))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.delete('/api/sloc-config/<sloc>')
def delete_sloc_config(sloc):
    try:
        SlocConfigService.delete_sloc(sloc)
        return '', 204
    except Exception as e:
        return jsonify({'error': str(e)}), 400

# ===== MATERIALS =====
@app.post('/api/materials')
def create_material():
    data = request.json

    try:
        MaterialService.create_or_update_material(
            material_number=data.get('material_number'),
            description=data.get('description'),
            base_uom=data.get('base_uom'),
            material_type=data.get('material_type'),
            material_group=data.get('material_group')
        )
        return jsonify(MaterialService.get_material(data.get('material_number'))), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/materials/<material_number>')
def get_material(material_number):
    material = MaterialService.get_material(material_number)
    if not material:
        return jsonify({'error': 'Material not found'}), 404
    return jsonify(material)

@app.get('/api/materials/search/<query>')
def search_materials(query):
    materials = MaterialService.search_materials(query)
    return jsonify(materials)

@app.get('/api/validate/material/<material_number>')
def validate_material(material_number):
    material = MaterialService.get_material(material_number)
    return jsonify({'exists': material is not None})

@app.get('/api/validate/wm-bin/<bin_code>')
def validate_wm_bin(bin_code):
    return jsonify({'exists': WmBinService.bin_exists(bin_code)})

@app.get('/api/validate/fixed-bin')
def validate_fixed_bin():
    material = request.args.get('material', '').strip()
    wm_bin = request.args.get('wm_bin', '').strip()
    if not material or not wm_bin:
        return jsonify({'error': 'material and wm_bin required'}), 400
    result = WmBinService.check_fixed_bin(material, wm_bin)
    return jsonify(result)

# ===== WM BINS =====
@app.post('/api/wm-bins')
def create_wm_bin():
    data = request.json

    try:
        bin_id = WmBinService.create_bin(
            bin_name=data.get('bin'),
            storage_type=data.get('storage_type'),
            sloc=data.get('sloc'),
            description=data.get('description')
        )
        return jsonify({'id': bin_id}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/wm-bins')
def list_wm_bins():
    sloc = request.args.get('sloc')
    storage_type = request.args.get('storage_type')

    if not sloc:
        return jsonify({'error': 'SLOC required'}), 400

    bins = WmBinService.list_bins(sloc, storage_type)
    return jsonify(bins)

@app.post('/api/wm-bins/<int:bin_id>/materials')
def add_material_to_bin(bin_id):
    data = request.json
    material_number = data.get('material_number')

    try:
        WmBinService.add_material_to_bin(bin_id, material_number)
        return '', 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/wm-bins/<int:bin_id>/materials')
def get_bin_materials(bin_id):
    materials = WmBinService.get_bin_materials(bin_id)
    return jsonify(materials)

@app.delete('/api/wm-bins/<int:bin_id>')
def delete_wm_bin(bin_id):
    try:
        WmBinService.delete_bin(bin_id)
        return '', 204
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.delete('/api/wm-bins/<int:bin_id>/materials/<material_number>')
def delete_bin_material(bin_id, material_number):
    try:
        WmBinService.delete_bin_material(bin_id, material_number)
        return '', 204
    except Exception as e:
        return jsonify({'error': str(e)}), 400

# ===== IMPORTS =====
def parse_csv_upload():
    file = request.files.get('file')
    if not file:
        return None, jsonify({'error': 'No file provided'}), 400

    filename = file.filename or ''
    rows = []

    # Handle XLSX files
    if filename.endswith('.xlsx'):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file.read()))
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(row):  # Skip empty rows
                    rows.append(dict(zip(headers, row)))
        except Exception as e:
            return None, jsonify({'error': f'Failed to read XLSX: {str(e)}'}), 400
    else:
        # Handle CSV and TXT files
        try:
            text = file.read().decode('utf-8-sig', errors='replace')
            reader = csv.DictReader(io.StringIO(text))
            rows = list(reader) if reader else []
        except Exception as e:
            return None, jsonify({'error': f'Failed to read file: {str(e)}'}), 400

    return rows, None, None

@app.get('/api/imports/status')
def get_import_status():
    return jsonify(ImportService.get_import_status())

@app.post('/api/imports/mara')
def import_mara():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_materials(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/makt')
def import_makt():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_materials(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/plant-data')
def import_plant_data():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_plant_data(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/valuation')
def import_valuation():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_valuation(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/snapshot')
def import_snapshot():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    session_id = request.form.get('sessionId')
    if not session_id:
        return jsonify({'error': 'sessionId required'}), 400
    try:
        count = ImportService.import_snapshot(int(session_id), rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/mard')
def import_mard():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_mard(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/mlgt')
def import_mlgt():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_mlgt(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/mlgn')
def import_mlgn():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_mlgn(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/lqua')
def import_lqua():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_lqua(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/storage-locations')
def import_storage_locations():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_storage_locations(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/lgap')
def import_lgap():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_lgap(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/mseg')
def import_mseg():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_mseg(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/lgplo')
def import_lgplo():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        count = ImportService.import_lgplo(rows)
        return jsonify({'imported': count}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/ekko')
def import_ekko():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        return jsonify({'imported': ImportService.import_ekko(rows)}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/ekpo')
def import_ekpo():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        return jsonify({'imported': ImportService.import_ekpo(rows)}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/aufk')
def import_aufk():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        return jsonify({'imported': ImportService.import_aufk(rows)}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/resb')
def import_resb():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        return jsonify({'imported': ImportService.import_resb(rows)}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.post('/api/imports/exclusions')
def import_exclusions():
    rows, err, code = parse_csv_upload()
    if err:
        return err, code
    try:
        return jsonify({'imported': ImportService.import_exclusions(rows)}), 201
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/imports/exclusions')
def list_exclusions():
    return jsonify(ImportService.get_excluded_materials())

@app.get('/api/validate/open-orders')
def validate_open_orders():
    material = request.args.get('material', '').strip()
    if not material:
        return jsonify({'error': 'material required'}), 400
    try:
        return jsonify(OrderValidationService.get_material_warnings(material))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sessions/<int:session_id>/counts-over-snapshot')
def counts_over_snapshot(session_id):
    try:
        return jsonify(OrderValidationService.get_counts_over_snapshot(session_id))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sessions/<int:session_id>/wm-discrepancies')
def wm_discrepancies(session_id):
    try:
        return jsonify(OrderValidationService.get_wm_discrepancies(session_id))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.get('/api/sessions/<int:session_id>/reservation-warnings')
def reservation_warnings(session_id):
    try:
        return jsonify(OrderValidationService.get_reservation_output_warnings(session_id))
    except Exception as e:
        return jsonify({'error': str(e)}), 400

# ===== EXPORT =====
@app.get('/api/sessions/<int:session_id>/export')
def export_session(session_id):
    fmt = request.args.get('format', 'csv')
    counts = CountService.list_counts(session_id, {})
    if not counts:
        counts = []

    output = io.StringIO()
    fields = ['id', 'material_number', 'quantity', 'sloc', 'wm_bin', 'zbin', 'status', 'username', 'created_at', 'updated_at']
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(counts)

    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': f'attachment; filename=session_{session_id}_counts.csv'}
    )

@app.get('/api/sessions/<int:session_id>/export/adjustments')
def export_adjustments(session_id):
    rows = ImportService.get_adjustment_export(session_id)
    output = io.StringIO()
    fields = ['material_number', 'description', 'plant', 'sloc', 'uom', 'snapshot_qty', 'counted_qty', 'adjustment']
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(rows)
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': f'attachment; filename=session_{session_id}_adjustments.csv'}
    )

# ===== AUDIT LOG =====
@app.get('/api/counts/<int:count_id>/audit')
def get_count_audit(count_id):
    logs = AuditService.get_audit_log(count_id)
    return jsonify(logs)

@app.get('/api/sessions/<int:session_id>/audit')
def get_session_audit(session_id):
    logs = AuditService.get_session_audit(session_id)
    return jsonify(logs)

# ===== DASHBOARD =====
@app.get('/api/sessions/<int:session_id>/dashboard/summary')
def get_dashboard_summary(session_id):
    summary = DashboardService.get_session_summary(session_id)
    return jsonify(summary)

@app.get('/api/sessions/<int:session_id>/dashboard/materials')
def get_material_summary(session_id):
    materials = DashboardService.get_material_summary(session_id)
    return jsonify(materials)

@app.get('/api/sessions/<int:session_id>/dashboard/users')
def get_user_summary(session_id):
    users = DashboardService.get_user_summary(session_id)
    return jsonify(users)

@app.get('/api/sessions/<int:session_id>/dashboard/discrepancies')
def get_discrepancies(session_id):
    discrepancies = DashboardService.get_discrepancies(session_id)
    return jsonify(discrepancies)

@app.get('/api/dashboard/problem-materials')
def get_problem_materials():
    data = DashboardService.get_problem_materials()
    return jsonify(data)

@app.get('/api/sessions/<int:session_id>/dashboard/high-value')
def get_high_value(session_id):
    data = DashboardService.get_high_value_materials(session_id)
    return jsonify(data)

# ===== HEALTH CHECK =====
@app.get('/health')
def health():
    return jsonify({'ok': True, 'time': datetime.now().isoformat()})

# ===== SERVE REACT FRONTEND =====
@app.route('/', defaults={'path': ''})
@app.route('/<path:path>')
def serve_frontend(path):
    if path.startswith('api/'):
        return jsonify({'error': 'Not found'}), 404

    # Serve actual static assets (JS, CSS, images, etc.)
    if path:
        asset = DIST_DIR / path
        if asset.exists() and asset.is_file():
            return send_from_directory(DIST_DIR, path)

    # All other paths (React routes) → serve index.html
    return send_from_directory(DIST_DIR, 'index.html')

@app.get('/cert')
def download_cert():
    cert_path = Path(__file__).parent / 'cert.pem'
    if cert_path.exists():
        return send_file(cert_path, as_attachment=True, download_name='physical-inventory.crt',
                         mimetype='application/x-x509-ca-cert')
    return jsonify({'error': 'No certificate. Run setup_ssl.py first.'}), 404

def _auto_init():
    """Ensure default seed data exists on every server start."""
    from init_db import init_seed_data
    import contextlib
    with contextlib.redirect_stdout(io.StringIO()):
        init_seed_data()

# Always run on startup (works with debug reloader too)
_auto_init()

if __name__ == '__main__':
    # Dev-only: direct Flask run on localhost.
    # For production (50 users), use: python run_production.py
    app.run(debug=True, host='127.0.0.1', port=8081)
